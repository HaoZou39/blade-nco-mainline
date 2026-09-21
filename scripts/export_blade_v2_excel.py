#!/usr/bin/env python
"""Export the current BLADE V2 run state to a formatted Excel workbook."""

from __future__ import annotations

from datetime import datetime
import json
import os
from pathlib import Path

import pandas as pd
import yaml
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
EVAL_ROOT = ROOT / "summary" / "evaluation_work"
OUTPUT = ROOT / "summary" / "BLADE_V2_M15_matched_results.xlsx"
TMP_OUTPUT = OUTPUT.with_suffix(".tmp.xlsx")
DATASETS = [
    "Uniform100",
    "GaussianMixture100",
    "TSPLIB100",
    "EvoReal_S1_100",
    "EvoReal_S2_100",
    "EvoReal_S3_100",
    "GaussianMixture500",
    "OriginalTSPLIB70",
]
PRIMARY_OOD = [
    "GaussianMixture100",
    "TSPLIB100",
    "EvoReal_S1_100",
    "EvoReal_S2_100",
    "EvoReal_S3_100",
]


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def model_index(value: str) -> int:
    return int(str(value).removeprefix("M"))


def evaluation_rows() -> pd.DataFrame:
    rows: list[dict] = []
    for path in EVAL_ROOT.rglob("summary.json"):
        value = read_json(path)
        rows.append(
            {
                "arm": value["arm"],
                "model": value["model"],
                "model_index": model_index(value["model"]),
                "dataset": value["dataset"],
                "mean_gap_percent": value["mean_gap_percent"],
                "median_gap_percent": value["median_gap_percent"],
                "P75": value["P75"],
                "P90": value["P90"],
                "P95": value["P95"],
                "P99": value["P99"],
                "max_gap_percent": value["max"],
                "valid_tours": value["valid_tours"],
                "failed_tours": value["failed_tours"],
                "instance_count": value["instance_count"],
                "inference_seconds": value["inference_wall_clock_sec"],
                "checkpoint_sha256": value["checkpoint_sha256"],
                "source": value["source"],
            }
        )
    frame = pd.DataFrame(rows)
    if len(frame):
        frame["dataset_order"] = frame.dataset.map({name: i for i, name in enumerate(DATASETS)})
        frame = frame.sort_values(["arm", "model_index", "dataset_order"]).drop(columns="dataset_order")
    return frame


def cycle_rows(arm: str, cycles: range) -> pd.DataFrame:
    run_id = "blade_v2_full_seed2026" if arm == "FULL" else "blade_v2_gp_only_seed2026"
    rows: list[dict] = []
    for cycle in cycles:
        path = ROOT / "runs" / run_id / "cycles" / f"C{cycle:02d}" / "cycle_summary.json"
        if not path.is_file():
            continue
        value = read_json(path)
        rows.append(
            {
                "arm": arm,
                "cycle": f"C{cycle:02d}",
                "cycle_index": cycle,
                "model_in": value["model_in"],
                "model_out": value["model_out"],
                "parent_policy": value["parent_policy"],
                "hardness_mean": value["hardness"]["mean"],
                "hardness_P95": value["hardness"]["P95"],
                "adaptation_mean_delta": value["adaptation_repair"]["mean_delta"],
                "adaptation_CI95_low": value["adaptation_repair"]["ci95_low"],
                "adaptation_CI95_high": value["adaptation_repair"]["ci95_high"],
                "fraction_improved": value["adaptation_repair"]["fraction_improved"],
                "parent_occupied_niches": value["parent_metrics"]["occupied_niches"],
                "parent_N_eff": value["parent_metrics"]["N_eff"],
                "parent_largest_niche_share": value["parent_metrics"]["largest_niche_share"],
                "offspring_occupied_niches": value["offspring_metrics"]["occupied_niches"],
                "offspring_N_eff": value["offspring_metrics"]["N_eff"],
                "archive_n": value["archive"]["archive_n"],
                "archive_new_unique": value["archive"]["new_unique"],
                "bq_epochs": value["bq"]["epochs"],
                "bq_learning_rate": value["bq"]["learning_rate"],
                "cycle_minutes": value["timing"]["cycle_wall_seconds"] / 60.0,
            }
        )
    return pd.DataFrame(rows)


def autosize_and_style(workbook) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column_cells in sheet.columns:
            letter = get_column_letter(column_cells[0].column)
            width = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
            sheet.column_dimensions[letter].width = min(max(width + 2, 10), 42)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"


def add_charts(workbook) -> None:
    trajectory = workbook["Full_Trajectory"]
    max_row = trajectory.max_row
    dataset_columns = {trajectory.cell(1, col).value: col for col in range(2, trajectory.max_column + 1)}

    high = LineChart()
    high.title = "FULL trajectory: high-shift datasets"
    high.y_axis.title = "Mean optimality gap (%)"
    high.x_axis.title = "Checkpoint"
    high.height = 8
    high.width = 15
    for name in ("GaussianMixture100", "GaussianMixture500", "EvoReal_S1_100"):
        col = dataset_columns[name]
        high.add_data(Reference(trajectory, min_col=col, min_row=1, max_row=max_row), titles_from_data=True)
    high.set_categories(Reference(trajectory, min_col=1, min_row=2, max_row=max_row))
    trajectory.add_chart(high, "K2")

    low = LineChart()
    low.title = "FULL trajectory: low-gap datasets"
    low.y_axis.title = "Mean optimality gap (%)"
    low.x_axis.title = "Checkpoint"
    low.height = 8
    low.width = 15
    for name in ("Uniform100", "TSPLIB100", "EvoReal_S2_100", "EvoReal_S3_100"):
        col = dataset_columns[name]
        low.add_data(Reference(trajectory, min_col=col, min_row=1, max_row=max_row), titles_from_data=True)
    low.set_categories(Reference(trajectory, min_col=1, min_row=2, max_row=max_row))
    trajectory.add_chart(low, "K18")

    comparison = workbook["M0_vs_M15"]
    bar = BarChart()
    bar.type = "col"
    bar.style = 10
    bar.title = "M0 vs M15 mean optimality gap"
    bar.y_axis.title = "Mean gap (%)"
    bar.height = 9
    bar.width = 16
    bar.add_data(Reference(comparison, min_col=2, max_col=3, min_row=1, max_row=comparison.max_row),
                 titles_from_data=True)
    bar.set_categories(Reference(comparison, min_col=1, min_row=2, max_row=comparison.max_row))
    comparison.add_chart(bar, "L2")


def main() -> None:
    config = yaml.safe_load((ROOT / "configs" / "blade_v2_protocol.yaml").read_text(encoding="utf-8-sig"))
    full_cycles_config = config["run"]["full_cycles"]
    gp_cycles_config = config["run"]["gp_only_cycles"]
    full_final_index = max(full_cycles_config) + 1
    gp_final_index = max(gp_cycles_config) + 1
    expected_evaluations = (full_final_index + 1 + gp_final_index - 1) * len(DATASETS)
    evaluations = evaluation_rows()
    full_eval = evaluations[evaluations.arm.eq("FULL")].copy()
    gp_eval = evaluations[evaluations.arm.eq("GP_ONLY")].copy()
    full_cycles = cycle_rows("FULL", range(min(full_cycles_config), max(full_cycles_config) + 1))
    gp_cycles = cycle_rows("GP_ONLY", range(min(gp_cycles_config), max(gp_cycles_config) + 1))

    m0 = full_eval[full_eval.model.eq("M0")].set_index("dataset")
    m15 = full_eval[full_eval.model.eq("M15")].set_index("dataset")
    comparison_rows = []
    for dataset in DATASETS:
        left = m0.loc[dataset]
        right = m15.loc[dataset]
        delta = right.mean_gap_percent - left.mean_gap_percent
        comparison_rows.append(
            {
                "dataset": dataset,
                "M0_mean_gap": left.mean_gap_percent,
                "M15_mean_gap": right.mean_gap_percent,
                "delta_pp_M15_minus_M0": delta,
                "relative_change_percent": 100.0 * delta / left.mean_gap_percent,
                "result": "IMPROVED" if delta < 0 else "WORSE",
                "M0_P95": left.P95,
                "M15_P95": right.P95,
                "M0_failed": left.failed_tours,
                "M15_failed": right.failed_tours,
            }
        )
    comparison = pd.DataFrame(comparison_rows)

    best = (
        full_eval.loc[full_eval.groupby("dataset").mean_gap_percent.idxmin(),
                      ["dataset", "model", "mean_gap_percent", "median_gap_percent", "P95", "failed_tours"]]
        .sort_values("dataset")
        .reset_index(drop=True)
    )

    trajectory = (
        full_eval.pivot(index="model_index", columns="dataset", values="mean_gap_percent")
        .reindex(columns=DATASETS)
        .reset_index()
        .rename(columns={"model_index": "model"})
    )
    trajectory["model"] = trajectory.model.map(lambda index: f"M{index}")

    shared_m1 = full_eval[full_eval.model.eq("M1")].copy()
    shared_m1["arm"] = "GP_ONLY"
    shared_m1["source"] = "shared M1 checkpoint; copied from FULL evaluation"
    ablation_eval = pd.concat(
        [
            full_eval[full_eval.model_index.between(1, gp_final_index)],
            shared_m1,
            gp_eval[gp_eval.model_index.between(2, gp_final_index)],
        ],
        ignore_index=True,
    ).sort_values(["model_index", "dataset", "arm"])

    m2 = ablation_eval[ablation_eval.model.eq("M2")]
    m2_pivot = m2.pivot(index="dataset", columns="arm", values="mean_gap_percent").reset_index()
    if {"FULL", "GP_ONLY"}.issubset(m2_pivot.columns):
        m2_pivot["GP_minus_FULL_pp"] = m2_pivot["GP_ONLY"] - m2_pivot["FULL"]
        m2_pivot["winner"] = m2_pivot["GP_minus_FULL_pp"].map(
            lambda value: "GP_ONLY" if value < 0 else "FULL"
        )
    matched_end = ablation_eval[ablation_eval.model.eq(f"M{gp_final_index}")]
    m15_pivot = matched_end.pivot(index="dataset", columns="arm", values="mean_gap_percent").reset_index()
    if {"FULL", "GP_ONLY"}.issubset(m15_pivot.columns):
        m15_pivot["GP_minus_FULL_pp"] = m15_pivot["GP_ONLY"] - m15_pivot["FULL"]
        m15_pivot["winner"] = m15_pivot["GP_minus_FULL_pp"].map(
            lambda value: "GP_ONLY" if value < 0 else "FULL"
        )

    ablation_trajectory = (
        ablation_eval.pivot(index=["model_index", "model", "arm"], columns="dataset",
                            values="mean_gap_percent")
        .reindex(columns=DATASETS)
        .reset_index()
        .sort_values(["model_index", "arm"])
    )

    progress_rows = []
    for arm, models in (("FULL", range(full_final_index + 1)),
                        ("GP_ONLY", range(2, gp_final_index + 1))):
        for index in models:
            subset = evaluations[evaluations.arm.eq(arm) & evaluations.model.eq(f"M{index}")]
            progress_rows.append(
                {
                    "arm": arm,
                    "model": f"M{index}",
                    "completed_datasets": len(subset),
                    "planned_datasets": 8,
                    "status": "COMPLETE" if len(subset) == 8 else "IN_PROGRESS" if len(subset) else "PENDING",
                }
            )
    progress = pd.DataFrame(progress_rows)

    primary_m0 = m0.loc[PRIMARY_OOD].mean_gap_percent.mean()
    primary_m15 = m15.loc[PRIMARY_OOD].mean_gap_percent.mean()
    heartbeat = read_json(ROOT / "runtime" / "heartbeat.json")
    eval_completed = len(evaluations)
    verification_path = ROOT / "summary" / "BLADE_V2_STRICT_VERIFICATION.json"
    verification = read_json(verification_path) if verification_path.is_file() else {}
    final_complete = (
        eval_completed == expected_evaluations
        and verification.get("status") == "PASS"
    )
    overview = pd.DataFrame(
        [
            ("Workbook status", "FINAL — all evaluations and strict verification passed"
             if final_complete else "PROVISIONAL — passive evaluation or verification is still running"),
            ("Generated at", datetime.now().astimezone().isoformat(timespec="seconds")),
            ("FULL training", f"{len(full_cycles)}/15 cycles complete; final model M15"),
            ("GP-only training", f"{len(gp_cycles)}/{len(gp_cycles_config)} cycles complete; final model M{gp_final_index}"),
            ("Passive evaluation", f"{eval_completed}/{expected_evaluations} summaries complete"),
            ("Current phase", f"{heartbeat.get('arm')} / {heartbeat.get('cycle')} / {heartbeat.get('phase')}"),
            ("Latest error", heartbeat.get("latest_error") or "None"),
            ("Primary N=100 OOD mean, M0", primary_m0),
            ("Primary N=100 OOD mean, M15", primary_m15),
            ("Primary OOD relative reduction", 100.0 * (primary_m0 - primary_m15) / primary_m0),
            ("Datasets improved M15 vs M0", f"{int((comparison.delta_pp_M15_minus_M0 < 0).sum())}/8"),
            ("Main trade-off", "Uniform100 worsened; shifted/OOD datasets generally improved"),
            ("Ablation status", f"FINAL — matched FULL/GP-only M1–M{gp_final_index} results included"
             if final_complete else f"Only completed GP-only model evaluations are included; final M{gp_final_index} conclusion pending"),
        ],
        columns=["item", "value"],
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(TMP_OUTPUT, engine="openpyxl") as writer:
        overview.to_excel(writer, sheet_name="Overview", index=False)
        comparison.to_excel(writer, sheet_name="M0_vs_M15", index=False)
        best.to_excel(writer, sheet_name="Best_Checkpoints", index=False)
        trajectory.to_excel(writer, sheet_name="Full_Trajectory", index=False)
        full_eval.to_excel(writer, sheet_name="Full_Evaluations", index=False)
        full_cycles.to_excel(writer, sheet_name="Full_Cycles", index=False)
        gp_cycles.to_excel(writer, sheet_name="GP_Cycles", index=False)
        ablation_eval.to_excel(writer, sheet_name="Ablation_Evals", index=False)
        ablation_trajectory.to_excel(writer, sheet_name="Ablation_Trajectory", index=False)
        m2_pivot.to_excel(writer, sheet_name="M2_Ablation", index=False)
        m15_pivot.to_excel(writer, sheet_name="M15_Ablation", index=False)
        progress.to_excel(writer, sheet_name="Eval_Progress", index=False)

    workbook = load_workbook(TMP_OUTPUT)
    autosize_and_style(workbook)
    workbook["Overview"].column_dimensions["A"].width = 34
    workbook["Overview"].column_dimensions["B"].width = 76
    workbook["M0_vs_M15"].conditional_formatting.add(
        f"D2:D{workbook['M0_vs_M15'].max_row}",
        ColorScaleRule(start_type="min", start_color="63BE7B",
                       mid_type="percentile", mid_value=50, mid_color="FFEB84",
                       end_type="max", end_color="F8696B"),
    )
    workbook["Full_Evaluations"].conditional_formatting.add(
        f"E2:E{workbook['Full_Evaluations'].max_row}",
        ColorScaleRule(start_type="min", start_color="63BE7B",
                       mid_type="percentile", mid_value=50, mid_color="FFEB84",
                       end_type="max", end_color="F8696B"),
    )
    add_charts(workbook)
    workbook.save(TMP_OUTPUT)
    os.replace(TMP_OUTPUT, OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
